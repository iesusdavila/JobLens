import io
from xml.sax.saxutils import escape
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Flowable, HRFlowable, Paragraph, SimpleDocTemplate, Spacer
from app.domain.cv_models import TailoredCv
from app.services.cv_templates.base import CvTemplate
from app.services.cv_templates.cv_blocks import CvBlock, CvBlockKind

class ClassicCvTemplate(CvTemplate):
    FONT_NAME = "Calibri"
    PDF_FONT = "Helvetica"
    PDF_BOLD_FONT = "Helvetica-Bold"

    @property
    def name(self) -> str:
        return "classic"

    def render_docx(self, cv: TailoredCv) -> bytes:
        document = Document()
        normal = document.styles["Normal"]
        normal.font.name = self.FONT_NAME
        normal.font.size = Pt(10.5)
        for block in self._block_builder.build(cv):
            self._add_docx_block(document, block)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def render_pdf(self, cv: TailoredCv) -> bytes:
        buffer = io.BytesIO()
        document = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=15 * mm, bottomMargin=15 * mm, title=cv.full_name or "CV")
        styles = self._pdf_styles()
        story: list[Flowable] = []
        for block in self._block_builder.build(cv):
            story.extend(self._pdf_flowables(block, styles))
        document.build(story)
        return buffer.getvalue()

    def _add_docx_block(self, document: Document, block: CvBlock) -> None:
        if block.kind == CvBlockKind.NAME:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block.text)
            run.bold = True
            run.font.size = Pt(18)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif block.kind in (CvBlockKind.CONTACT, CvBlockKind.HEADLINE):
            paragraph = document.add_paragraph(block.text)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif block.kind == CvBlockKind.HEADING:
            document.add_heading(block.text, level=1)
        elif block.kind == CvBlockKind.ENTRY:
            document.add_paragraph().add_run(block.text).bold = True
        elif block.kind == CvBlockKind.ENTRY_META:
            document.add_paragraph().add_run(block.text).italic = True
        elif block.kind == CvBlockKind.BULLET:
            document.add_paragraph(block.text, style="List Bullet")
        else:
            document.add_paragraph(block.text)

    def _pdf_styles(self) -> dict[CvBlockKind, ParagraphStyle]:
        base = ParagraphStyle("base", fontName=self.PDF_FONT, fontSize=10, leading=13.5)
        return {
            CvBlockKind.NAME: ParagraphStyle("name", parent=base, fontName=self.PDF_BOLD_FONT, fontSize=18, leading=22, alignment=TA_CENTER),
            CvBlockKind.HEADLINE: ParagraphStyle("headline", parent=base, fontSize=11, alignment=TA_CENTER),
            CvBlockKind.CONTACT: ParagraphStyle("contact", parent=base, fontSize=9, alignment=TA_CENTER, spaceAfter=6),
            CvBlockKind.HEADING: ParagraphStyle("heading", parent=base, fontName=self.PDF_BOLD_FONT, fontSize=12, leading=15, spaceBefore=8),
            CvBlockKind.ENTRY: ParagraphStyle("entry", parent=base, fontName=self.PDF_BOLD_FONT, spaceBefore=4),
            CvBlockKind.ENTRY_META: ParagraphStyle("meta", parent=base, fontName="Helvetica-Oblique", fontSize=9),
            CvBlockKind.BULLET: ParagraphStyle("bullet", parent=base, leftIndent=12, bulletIndent=2),
            CvBlockKind.PARAGRAPH: base,
        }

    def _pdf_flowables(self, block: CvBlock, styles: dict[CvBlockKind, ParagraphStyle]) -> list[Flowable]:
        text = escape(block.text)
        if block.kind == CvBlockKind.HEADING:
            return [Paragraph(text, styles[block.kind]), HRFlowable(width="100%", thickness=0.6, spaceBefore=1, spaceAfter=3)]
        if block.kind == CvBlockKind.BULLET:
            return [Paragraph(text, styles[block.kind], bulletText="•")]
        if block.kind == CvBlockKind.CONTACT:
            return [Paragraph(text, styles[block.kind]), Spacer(1, 2)]
        return [Paragraph(text, styles[block.kind])]
