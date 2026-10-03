import io
import re
from abc import ABC, abstractmethod
from pathlib import Path
from zipfile import BadZipFile
from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from app.core.exceptions import DocumentParsingError, FileTooLargeError, UnsupportedFileTypeError
from app.domain.document_models import DocumentSection, ParsedDocument
from app.domain.enums import DocumentKind

class TextNormalizer:
    def normalize(self, text: str) -> str:
        text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ").replace("\t", " ")
        text = re.sub(r"[​-‍﻿]", "", text)
        lines = [re.sub(r" {2,}", " ", line).strip() for line in text.split("\n")]
        text = "\n".join(lines)
        return re.sub(r"\n{3,}", "\n\n", text).strip()

class SectionSplitter:
    KNOWN_HEADINGS = {
        "summary", "profile", "professional summary", "about me", "objective", "experience", "work experience",
        "professional experience", "employment history", "education", "skills", "technical skills", "projects",
        "certifications", "certificates", "languages", "publications", "awards", "achievements", "volunteering",
        "resumen", "perfil", "experiencia", "experiencia laboral", "experiencia profesional", "educación", "educacion",
        "formación", "formacion", "habilidades", "competencias", "proyectos", "certificaciones", "idiomas", "logros",
    }
    MAX_HEADING_WORDS = 5

    def split(self, text: str, kind: DocumentKind) -> list[DocumentSection]:
        sections: list[DocumentSection] = []
        heading = "Header"
        buffer: list[str] = []
        for line in text.split("\n"):
            if self._is_heading(line):
                self._flush(sections, kind, heading, buffer)
                heading = line.strip("#: ").strip()
                buffer = []
                continue
            buffer.append(line)
        self._flush(sections, kind, heading, buffer)
        return sections

    def _is_heading(self, line: str) -> bool:
        stripped = line.strip()
        if not stripped or len(stripped.split()) > self.MAX_HEADING_WORDS:
            return False
        normalized = stripped.strip("#: ").strip().lower()
        if stripped.startswith("#") or normalized in self.KNOWN_HEADINGS:
            return True
        letters = [char for char in stripped if char.isalpha()]
        return len(letters) >= 4 and all(char.isupper() for char in letters)

    @staticmethod
    def _flush(sections: list[DocumentSection], kind: DocumentKind, heading: str, buffer: list[str]) -> None:
        body = "\n".join(buffer).strip()
        if not body:
            return
        sections.append(DocumentSection(id=f"{kind.value}-s{len(sections) + 1}", source=kind, heading=heading, text=body))

class DocumentParser(ABC):
    @property
    @abstractmethod
    def extensions(self) -> tuple[str, ...]:
        raise NotImplementedError

    @abstractmethod
    def extract_text(self, content: bytes) -> str:
        raise NotImplementedError

    def matches_signature(self, content: bytes) -> bool:
        return True

class PdfDocumentParser(DocumentParser):
    @property
    def extensions(self) -> tuple[str, ...]:
        return (".pdf",)

    def matches_signature(self, content: bytes) -> bool:
        return content.lstrip()[:5] == b"%PDF-"

    def extract_text(self, content: bytes) -> str:
        try:
            reader = PdfReader(io.BytesIO(content))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except (PdfReadError, ValueError, KeyError) as error:
            raise DocumentParsingError("The PDF file is corrupted or encrypted. Export it again or paste the text.") from error

class DocxDocumentParser(DocumentParser):
    @property
    def extensions(self) -> tuple[str, ...]:
        return (".docx",)

    def matches_signature(self, content: bytes) -> bool:
        return content[:2] == b"PK"

    def extract_text(self, content: bytes) -> str:
        try:
            document = DocxDocument(io.BytesIO(content))
        except (PackageNotFoundError, BadZipFile, KeyError, ValueError) as error:
            raise DocumentParsingError("The DOCX file could not be opened. Save it again as .docx or paste the text.") from error
        lines = [self._paragraph_text(paragraph) for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
        return "\n".join(lines)

    @staticmethod
    def _paragraph_text(paragraph: Paragraph) -> str:
        text = paragraph.text.strip()
        style_name = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        if text and style_name.startswith(("heading", "title")):
            return f"# {text}"
        if text and "list" in style_name:
            return f"- {text}"
        return text

class PlainTextDocumentParser(DocumentParser):
    ENCODINGS = ("utf-8", "utf-16", "latin-1")

    @property
    def extensions(self) -> tuple[str, ...]:
        return (".txt", ".md", ".markdown")

    def matches_signature(self, content: bytes) -> bool:
        return b"\x00" not in content[:1024] or content[:2] in (b"\xff\xfe", b"\xfe\xff")

    def extract_text(self, content: bytes) -> str:
        for encoding in self.ENCODINGS:
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise DocumentParsingError("The text file encoding is not supported. Save it as UTF-8.")

class DocumentParserService:
    MIN_TEXT_CHARS = 50

    def __init__(
        self,
        parsers: list[DocumentParser],
        normalizer: TextNormalizer,
        splitter: SectionSplitter,
        max_upload_bytes: int,
    ) -> None:
        self._parsers = {extension: parser for parser in parsers for extension in parser.extensions}
        self._normalizer = normalizer
        self._splitter = splitter
        self._max_upload_bytes = max_upload_bytes

    @property
    def supported_extensions(self) -> list[str]:
        return sorted(self._parsers)

    def parse_file(self, filename: str, content: bytes, kind: DocumentKind) -> ParsedDocument:
        if len(content) > self._max_upload_bytes:
            raise FileTooLargeError(f"'{filename}' exceeds the {self._max_upload_bytes // (1024 * 1024)} MB limit.")
        parser = self._parser_for(filename)
        if not parser.matches_signature(content):
            raise UnsupportedFileTypeError(f"'{filename}' content does not match its extension.")
        return self._build(filename, parser.extract_text(content), kind)

    def parse_text(self, text: str, kind: DocumentKind) -> ParsedDocument:
        return self._build(f"pasted-{kind.value}.txt", text, kind)

    def _parser_for(self, filename: str) -> DocumentParser:
        extension = Path(filename).suffix.lower()
        parser = self._parsers.get(extension)
        if parser is None:
            raise UnsupportedFileTypeError(f"Unsupported file type '{extension or 'unknown'}'. Use one of: {', '.join(self.supported_extensions)}.")
        return parser

    def _build(self, filename: str, raw_text: str, kind: DocumentKind) -> ParsedDocument:
        text = self._normalizer.normalize(raw_text)
        if len(text) < self.MIN_TEXT_CHARS:
            raise DocumentParsingError(
                f"Almost no text could be read from '{filename}'. If it is a scanned PDF or an image, paste the text instead."
            )
        return ParsedDocument(kind=kind, filename=filename, text=text, sections=self._splitter.split(text, kind))
