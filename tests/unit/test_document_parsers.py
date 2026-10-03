import io
import pytest
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from app.core.exceptions import DocumentParsingError, FileTooLargeError, UnsupportedFileTypeError
from app.domain.enums import DocumentKind
from app.services.document_parser_service import DocumentParserService, DocxDocumentParser, PdfDocumentParser, PlainTextDocumentParser, SectionSplitter, TextNormalizer
from tests.sample_data import CV_TEXT

def service(max_bytes: int = 1024 * 1024) -> DocumentParserService:
    return DocumentParserService([PdfDocumentParser(), DocxDocumentParser(), PlainTextDocumentParser()], TextNormalizer(), SectionSplitter(), max_bytes)

def pdf_bytes(lines: list[str]) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    y = 800
    for line in lines:
        pdf.drawString(50, y, line)
        y -= 16
    pdf.save()
    return buffer.getvalue()

def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_heading("Experience", level=1)
    document.add_paragraph("Senior Backend Engineer at Acme Corp building FastAPI services for millions of users.")
    document.add_heading("Skills", level=1)
    document.add_paragraph("Python", style="List Bullet")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()

def test_plain_text_is_normalized_and_split_into_sections() -> None:
    parsed = service().parse_file("cv.txt", ("  " + CV_TEXT.replace("\n", "\r\n") + "\n\n\n\n").encode("utf-8"), DocumentKind.CV)
    headings = [section.heading for section in parsed.sections]
    assert headings == ["Header", "SUMMARY", "EXPERIENCE", "EDUCATION", "SKILLS"]
    assert "\r" not in parsed.text and "\n\n\n" not in parsed.text
    assert parsed.sections[2].id == "cv-s3"

def test_docx_headings_and_bullets_are_preserved() -> None:
    parsed = service().parse_file("cv.docx", docx_bytes(), DocumentKind.CV)
    assert [section.heading for section in parsed.sections] == ["Header", "Experience", "Skills"]
    assert "- Python" in parsed.text

def test_pdf_text_is_extracted() -> None:
    parsed = service().parse_file("cv.pdf", pdf_bytes(["Jane Doe", "EXPERIENCE", "Senior Backend Engineer at Acme Corp with FastAPI"]), DocumentKind.CV)
    assert "Acme Corp" in parsed.text
    assert "EXPERIENCE" in [section.heading for section in parsed.sections]

def test_scanned_pdf_without_text_asks_for_paste() -> None:
    with pytest.raises(DocumentParsingError, match="paste the text"):
        service().parse_file("scan.pdf", pdf_bytes([]), DocumentKind.CV)

def test_unsupported_extension_is_rejected() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        service().parse_file("cv.exe", b"MZ....", DocumentKind.CV)

def test_mismatched_signature_is_rejected() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        service().parse_file("cv.pdf", b"not really a pdf", DocumentKind.CV)

def test_large_file_is_rejected() -> None:
    with pytest.raises(FileTooLargeError):
        service(max_bytes=10).parse_file("cv.txt", b"x" * 100, DocumentKind.CV)

def test_pasted_text_is_parsed() -> None:
    parsed = service().parse_text(CV_TEXT, DocumentKind.EXTRA)
    assert parsed.kind == DocumentKind.EXTRA
    assert parsed.sections[0].id == "extra-s1"
